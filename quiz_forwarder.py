import os
import sys
import re
import time
import logging
from datetime import datetime
import pandas as pd
import openpyxl
from playwright.sync_api import sync_playwright

# ================= SYSTEM & LOGGING SETUP =================
BASE_DIR = r"C:\EnglishCourseAutomation"
os.makedirs(BASE_DIR, exist_ok=True)

LOG_FILE = os.path.join(BASE_DIR, "quiz_automation.log")
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

# ================= CONFIGURATION =================
EXCEL_FILENAME = "Hashims Community Group.xlsx"
SHEET_NAME = "Fluence(09Sep)"

# Exact name of your shared staging group:
SOURCE_CHAT_NAME = "Community automation"

PROFILE_DIR = os.path.join(BASE_DIR, "whatsapp_profile")
BATCH_SIZE = 5

# Constant footer line on the daily video post
VIDEO_ANCHOR_LINE = "Do the Shadow Writing"

# Set to None to automatically use today's date (1 to 31)
TARGET_DAY = None
MARK_TEXT = "SWAP"

# Weekly send limits (checked across the last 7 day columns in Excel)
WEEKLY_LIMITS = {
    "test": None,
    "hot": None,
    "funnel 2": None,
    "funnel2": None,
    "funnel  2": None,
    "warm": 2,          
    "fluencer": None,
    "fluvencer": None,
    "cool": 1,          
    "a_fluencer": None,
    "staff fluencer": None,
    "staff_fluencer": None,
}
# =================================================


def resolve_excel_path():
    exe_dir = os.path.dirname(os.path.abspath(sys.argv[0]))
    candidate_paths = [
        os.path.join(exe_dir, EXCEL_FILENAME),
        os.path.join(os.getcwd(), EXCEL_FILENAME),
        os.path.join(BASE_DIR, EXCEL_FILENAME)
    ]
    for path in candidate_paths:
        if os.path.exists(path):
            return path
    return os.path.join(BASE_DIR, EXCEL_FILENAME)


def get_target_day():
    return TARGET_DAY if TARGET_DAY is not None else datetime.now().day


def _find_day_col_name(df_columns, d):
    if d in df_columns:
        return d
    if str(d) in df_columns:
        return str(d)
    return None


def count_recent_sends(row, df_columns, current_day, window_days=7):
    start_day = max(1, current_day - (window_days - 1))
    count = 0
    for d in range(start_day, current_day + 1):
        col = _find_day_col_name(df_columns, d)
        if col is not None:
            cell_val = str(row.get(col, "")).strip().upper()
            if MARK_TEXT.upper() in cell_val or "SWCP" in cell_val:
                count += 1
    return count


def load_and_sort_contacts(excel_path, sheet_name, day_num):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(
            f"Could not find '{EXCEL_FILENAME}'. Please place it in {BASE_DIR} or next to QuizForwarder.exe."
        )

    print(f"📊 Reading Excel file: {excel_path} (Sheet: {sheet_name}, Day Column: {day_num})...")
    df = pd.read_excel(excel_path, sheet_name=sheet_name)

    df["Excel_Row"] = df.index + 2
    df["Type_Clean"] = df["Type"].astype(str).str.strip().str.lower().str.replace("  ", " ")
    df["Funnel_Clean"] = df.get("Funnel 2 State", pd.Series("", index=df.index)).astype(str).str.strip().str.lower()

    test_rows = df[(df["Type_Clean"] == "test") | (df["Funnel_Clean"] == "test")].copy()
    if len(test_rows) > 0:
        print("\n" + "🧪" * 28)
        print(f"🧪 TEST MODE ACTIVE: Found {len(test_rows)} row(s) marked 'Test'!")
        print("🧪 Only sending to 'Test' contacts (All other contacts are safely ignored).")
        print("🧪" * 28 + "\n")
        valid_df = test_rows.reset_index(drop=True)
        valid_df["Priority"] = 0
        is_test_mode = True
    else:
        is_test_mode = False

        def assign_priority(type_val):
            type_str = str(type_val).lower()
            if "hot" in type_str or ("funnel" in type_str and "2" in type_str):
                return 1
            elif type_str in ["warm", "fluencer", "fluvencer"]:
                return 2
            elif type_str in ["cool", "a_fluencer", "staff fluencer", "staff_fluencer"]:
                return 3
            return 99

        df["Priority"] = df["Type_Clean"].apply(assign_priority)
        valid_df = df[df["Priority"] < 99].copy()
        valid_df = valid_df.sort_values(by="Priority", kind="mergesort").reset_index(drop=True)

    day_col = _find_day_col_name(valid_df.columns, day_num)
    if day_col is None:
        raise Exception(f"Could not find day column '{day_num}' in sheet '{sheet_name}'.")

    contacts = []
    skipped_today = 0
    skipped_weekly_cap = 0

    for _, row in valid_df.iterrows():
        type_clean = row["Type_Clean"]
        current_mark = str(row.get(day_col, "")).strip()

        if not is_test_mode and MARK_TEXT.lower() in current_mark.lower():
            skipped_today += 1
            continue

        weekly_cap = WEEKLY_LIMITS.get(type_clean, None)
        if "funnel" in type_clean and "2" in type_clean:
            weekly_cap = None
            
        if not is_test_mode and weekly_cap is not None:
            sends_this_week = count_recent_sends(row, valid_df.columns, day_num, window_days=7)
            if sends_this_week >= weekly_cap:
                skipped_weekly_cap += 1
                continue

        raw_name = str(row.get("Name", "")).strip()
        raw_num = str(row.get("Number", "")).strip()

        first_num_line = raw_num.split("\n")[0].strip()
        digits_only = re.sub(r"\D", "", first_num_line)

        contacts.append({
            "excel_row": int(row["Excel_Row"]),
            "name": raw_name,
            "raw_number": first_num_line,
            "digits": digits_only,
            "type": str(row.get("Type", "")).strip(),
            "priority": row["Priority"]
        })

    print(
        f"✅ Loaded {len(contacts)} eligible contact(s) for Day {day_num} "
        f"(Skipped {skipped_today} already marked '{MARK_TEXT}' today, "
        f"and {skipped_weekly_cap} at their Warm/Cool weekly limit)."
    )
    return contacts


def mark_contacts_as_swap(excel_path, sheet_name, day_num, completed_contacts):
    if not completed_contacts:
        return

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb[sheet_name]

        target_col_idx = None
        for col_idx in range(1, ws.max_column + 1):
            cell_val = ws.cell(row=1, column=col_idx).value
            if str(cell_val).strip() == str(day_num):
                target_col_idx = col_idx
                break

        if target_col_idx is None:
            print(f"⚠️ Could not find Day {day_num} header in Excel to mark '{MARK_TEXT}'.")
            return

        for contact in completed_contacts:
            row_idx = contact["excel_row"]
            existing = ws.cell(row=row_idx, column=target_col_idx).value
            if existing and str(existing).strip():
                if MARK_TEXT.lower() not in str(existing).lower():
                    ws.cell(row=row_idx, column=target_col_idx).value = f"{existing}\n{MARK_TEXT}"
            else:
                ws.cell(row=row_idx, column=target_col_idx).value = MARK_TEXT

        wb.save(excel_path)
        wb.close()
        print(f"💾 Marked '{MARK_TEXT}' in Day {day_num} column for {len(completed_contacts)} contact(s).")
    except PermissionError:
        print("❌ Could not save Excel file! Please close 'Hashims Community Group.xlsx' in Excel so Python can update it.")
    except Exception as e:
        print(f"⚠️ Error saving '{MARK_TEXT}' to Excel: {e}")


class WhatsAppQuizForwarder:
    def __init__(self, profile_dir):
        self.profile_dir = profile_dir
        os.makedirs(self.profile_dir, exist_ok=True)
        self.playwright = sync_playwright().start()

        try:
            self.browser = self.playwright.chromium.launch_persistent_context(
                user_data_dir=self.profile_dir,
                headless=False,
                channel="chrome",
                args=["--start-maximized"]
            )
        except Exception:
            self.browser = self.playwright.chromium.launch_persistent_context(
                user_data_dir=self.profile_dir,
                headless=False,
                channel="msedge",
                args=["--start-maximized"]
            )

        self.page = self.browser.pages[0]
        self.page.goto("https://web.whatsapp.com/", timeout=0)
        print("Waiting for WhatsApp Web to load...")

        try:
            self.page.wait_for_selector("#pane-side, header", timeout=20000)
        except Exception:
            print("Please scan the WhatsApp QR code if prompted. Waiting up to 2 minutes...")
            self.page.wait_for_selector("#pane-side, header", timeout=120000)

        print("WhatsApp Web is ready.")
        time.sleep(2)

    def _force_scroll_to_very_bottom(self):
        for _ in range(2):
            self.page.mouse.move(850, 500)
            self.page.mouse.wheel(0, 4000)
            self.page.evaluate("""() => {
                const main = document.querySelector('#main');
                if (!main) return;
                const divs = Array.from(main.querySelectorAll('div'));
                for (const d of divs) {
                    if (d.scrollHeight > d.clientHeight + 50) {
                        d.scrollTop = d.scrollHeight;
                    }
                }
            }""")
            time.sleep(0.4)

    def _is_chat_actually_open(self, target_name):
        """Checks if the right-hand #main chat pane is open to the correct group."""
        return self.page.evaluate("""(name) => {
            const header = document.querySelector('#main header');
            if (!header) return false;
            
            // Read the first line of the header text (which is always the group title)
            const hText = (header.innerText || '').trim();
            const firstLine = hText.split('\\n')[0].toLowerCase();
            
            // Check if the group name is explicitly found in that first line
            if (!firstLine.includes(name.toLowerCase())) return false;
            
            const msgCount = document.querySelectorAll('#main .message-in, #main .message-out, #main [data-id], #main [role="row"]').length;
            return msgCount > 0;
        }""", target_name)

    def open_source_chat(self, chat_name):
        time.sleep(2.0)

        if self._is_chat_actually_open(chat_name):
            return

        print(f"🔍 Navigating back to source group: '{chat_name}'...")

        search_box = self.page.locator('#side [contenteditable="true"], #side input[type="text"], #side [role="textbox"]').first
        if search_box.is_visible():
            search_box.click()
            self.page.keyboard.press("Control+A")
            self.page.keyboard.press("Backspace")
            time.sleep(0.5)

            self.page.keyboard.insert_text(chat_name)
            time.sleep(2.0)

            matched_row = self.page.locator("#pane-side [role='row']").get_by_text(chat_name, exact=False).first
            if matched_row.is_visible():
                matched_row.click(force=True)
            else:
                self.page.keyboard.press("ArrowDown")
                time.sleep(0.3)
                self.page.keyboard.press("Enter")

            time.sleep(2.5)
            
            if self._is_chat_actually_open(chat_name):
                print(f"✅ Re-opened '{chat_name}' successfully!")
                return
            else:
                search_box.click()
                self.page.keyboard.press("Control+A")
                self.page.keyboard.press("Backspace")
                time.sleep(1.0)

        print("\n" + "=" * 60)
        print(f"👉 Please manually click on '{chat_name}' in the left chat list of WhatsApp Web!")
        print("=" * 60)
        for _ in range(60):
            if self._is_chat_actually_open(chat_name):
                print(f"✅ '{chat_name}' is now explicitly verified! Resuming...")
                time.sleep(1)
                return
            time.sleep(2)

        raise Exception(f"Could not return to '{chat_name}'. Script aborted to prevent sending wrong messages.")

    def _get_forward_modal_and_search(self):
        dialog_inputs = self.page.locator(
            'div[role="dialog"] div[contenteditable="true"], '
            'div[role="dialog"] input[type="text"], '
            '[data-animate-modal-popup="true"] div[contenteditable="true"], '
            '[data-animate-modal-popup="true"] input[type="text"]'
        ).all()
        for inp in dialog_inputs:
            if inp.is_visible():
                return inp

        has_forward_title = self.page.evaluate("""() => {
            const bodyTxt = document.body.innerText || '';
            return bodyTxt.includes('Forward message to') ||
                   bodyTxt.includes('Forward update to') ||
                   bodyTxt.includes('Forward messages to');
        }""")
        if has_forward_title:
            all_inputs = self.page.locator('div[contenteditable="true"], input[type="text"]').all()
            for inp in all_inputs:
                if not inp.is_visible():
                    continue
                box = inp.bounding_box()
                if box and 300 < box["x"] < 700 and 140 < box["y"] < 360:
                    return inp

        return None

    def _close_modal_if_open(self):
        if self._get_forward_modal_and_search() is not None:
            close_btn = self.page.locator(
                'div[role="dialog"] [aria-label="Close"], '
                'div[role="dialog"] [data-icon="x"], '
                '[data-animate-modal-popup="true"] [data-icon="x"]'
            ).first
            if close_btn.is_visible():
                close_btn.click(force=True)
            else:
                self.page.keyboard.press("Escape")
            time.sleep(0.6)

    def _check_and_advance_to_modal(self):
        if self._get_forward_modal_and_search() is not None:
            return True

        fwd_menu = self.page.locator(
            'li:has-text("Forward"), '
            'div[role="button"]:has-text("Forward"), '
            '[data-animate-dropdown-item="true"]:has-
